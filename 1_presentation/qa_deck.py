"""
QA of the final deck, measured by PowerPoint itself (COM), then PDF export and PDF checks.

Checks per slide: text overflow (rendered text bounds vs box), content overlaps (text / tables / our figures; template
artwork excluded), off-slide content, smallest font, picture resolution, one status pill per content slide.
Deck checks: slide count, section order and slide limits vs the official template, fonts used and embedded,
PDF page count/size/fonts. Writes qa_results.json (next to this script), the PDF, and PNG renders into a scratch folder.
Usage: python qa_deck.py [render_dir]
"""
import json
import sys
from pathlib import Path

import pythoncom
import win32com.client

HERE = Path(__file__).resolve().parent
DECK = HERE.parent / "FINAL_SUBMISSION" / "Vishwakarma_Stage2_Final.pptx"
PDF = HERE.parent / "FINAL_SUBMISSION" / "Vishwakarma_Stage2_Final.pdf"
OUT = HERE / "qa_results.json"            # working evidence, kept outside the deliverables folder
RENDER = Path(sys.argv[1]) if len(sys.argv) > 1 else None

# official template sections in order, with their slide limits
SECTIONS = [("Application ID (cover)", ("Vishwakarma Application ID",), 1, 1),
            ("Problem Context", ("Problem Context", "Project Overview"), 1, 2),   # T2: Project Overview + Problem Statement
            ("User Journey", ("User Journey",), 1, 2),
            ("System Architecture & Specs", ("System Architecture & Specs",), 1, 2),
            ("CAD Design & Mechanical Specs", ("CAD Design & Mechanical Specs",), 1, 2),
            ("Electronics Schematics & BOM", ("Electronics Schematics & BOM",), 1, 2),
            ("Maker Journey", ("Maker Journey",), 1, 2),
            ("Safety & Hazard Mitigation", ("Safety & Hazard Mitigation",), 1, 1),
            ("Roadmap", ("Roadmap",), 1, 1),
            ("Technical Challenges & Risks", ("Technical Challenges & Risks",), 1, 1),
            ("Potential Solutions & Risk Mitigation", ("Potential Solutions & Risk Mitigation",), 1, 1),
            ("References & Sources", ("References & Sources",), 1, 1),
            ("Image Sources", ("Image Sources",), 1, 1)]
PT = 72.0


def rect(sh):
    return sh.Left, sh.Top, sh.Left + sh.Width, sh.Top + sh.Height


def text_rect(sh):
    tr = sh.TextFrame2.TextRange
    return tr.BoundLeft, tr.BoundTop, tr.BoundLeft + tr.BoundWidth, tr.BoundTop + tr.BoundHeight


def overlap(a, b):
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def main():
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    res = {"slides": [], "deck": {}}
    try:
        pres = app.Presentations.Open(str(DECK), True, False, False)
        W, H = pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight
        res["deck"]["slide_count"] = pres.Slides.Count
        res["deck"]["slide_size_in"] = (round(W / PT, 3), round(H / PT, 3))
        res["deck"]["fonts"] = sorted({(pres.Fonts(i).Name, bool(pres.Fonts(i).Embedded)) for i in range(1, pres.Fonts.Count + 1)})
        titles = []
        for i in range(1, pres.Slides.Count + 1):
            s = pres.Slides(i)
            issues, content, min_font, pills = [], [], 99.0, 0
            for j in range(1, s.Shapes.Count + 1):
                sh = s.Shapes(j)
                ours = not sh.Name.startswith("Google Shape")
                r = rect(sh)
                if sh.HasTable:
                    content.append((sh.Name, r, "table"))
                    continue
                if sh.Type == 13 and ours:                         # our pictures (figures)
                    content.append((sh.Name, r, "picture"))
                    continue
                if sh.HasTextFrame and sh.TextFrame2.HasText:
                    t = sh.TextFrame2.TextRange
                    if sh.Top < 45 and sh.Left > 600 and sh.AutoShapeType == 5:
                        pills += 1
                    tb = text_rect(sh)
                    if (sh.Top < 45 and t.Font.Size >= 18) or t.Text.strip() == "Vishwakarma Application ID":
                        titles.append((i, t.Text.strip()))
                    if tb[3] > r[3] + 3 or tb[2] > r[2] + 3:
                        issues.append(f"text overflows its box: '{t.Text[:50]}' (text bottom {tb[3] / PT:.2f} in vs box "
                                      f"{r[3] / PT:.2f} in)")
                    if tb[0] < -1 or tb[1] < -1 or tb[2] > W + 1 or tb[3] > H + 1:
                        issues.append(f"text off the slide: '{t.Text[:50]}'")
                    content.append((sh.Name, tb, "text:" + t.Text[:40].replace("\n", " ")))
            for a in range(len(content)):
                for b in range(a + 1, len(content)):
                    ov = overlap(content[a][1], content[b][1])
                    if ov > 25:                                    # > 25 pt² (about 0.005 in²)
                        issues.append(f"overlap {ov:.0f} pt²: [{content[a][2]}] x [{content[b][2]}]")
            res["slides"].append({"n": i, "issues": issues, "min_font_pt": round(min_font, 2), "status_pills": pills,
                                  "content_shapes": len(content)})
        res["deck"]["titles"] = titles
        pres.SaveAs(str(PDF), 32)                                  # ppSaveAsPDF
        if RENDER:
            RENDER.mkdir(parents=True, exist_ok=True)
            for i in range(1, pres.Slides.Count + 1):
                pres.Slides(i).Export(str(RENDER / f"s{i:02d}.png"), "PNG", 1600, 900)
        pres.Close()
    finally:
        app.Quit()

    # sections: order and limits from the titles
    seq = []
    for n, t in res["deck"]["titles"]:
        for name, keys, lo, hi in SECTIONS:
            if any(t.startswith(k) for k in keys):
                seq.append((n, name))
                break
        else:
            seq.append((n, f"UNMAPPED: {t}"))
    order_ok, counts, last = True, {}, -1
    for n, name in seq:
        idx = next((k for k, (nm, *_rest) in enumerate(SECTIONS) if nm == name), None)
        if idx is None or idx < last:
            order_ok = False
        last = max(last, idx if idx is not None else last)
        counts[name] = counts.get(name, 0) + 1
    limits = {name: (counts.get(name, 0), lo, hi, lo <= counts.get(name, 0) <= hi) for name, _, lo, hi in SECTIONS}
    res["deck"]["section_sequence"] = seq
    res["deck"]["section_order_ok"] = order_ok
    res["deck"]["section_limits"] = limits

    # smallest font per slide from the file itself (python-pptx), table cells included
    from pptx import Presentation
    from pptx.util import Pt
    for i, s in enumerate(Presentation(DECK).slides, 1):
        sizes = []
        for sh in s.shapes:
            frames = [sh.text_frame] if sh.has_text_frame else []
            if sh.has_table:
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
            for tf in frames:
                sizes += [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size and r.text.strip()]
        res["slides"][i - 1]["min_font_pt"] = min(sizes) if sizes else None
    low = []
    for i, s in enumerate(Presentation(DECK).slides, 1):
        for sh in s.shapes:
            if sh.shape_type == 13 and not sh.name.startswith("Google Shape"):
                px_w = sh.image.size[0]
                dpi = px_w / (sh.width / 914400)
                if dpi < 150:
                    low.append((i, sh.name, round(dpi)))
    res["deck"]["low_res_pictures"] = low

    # PDF checks
    import pymupdf
    d = pymupdf.open(PDF)
    fonts = sorted({f[3] for p in d for f in p.get_fonts()})
    res["deck"]["pdf"] = {"pages": d.page_count, "size_in": (round(d[0].rect.width / 72, 2), round(d[0].rect.height / 72, 2)),
                          "fonts": fonts, "titles_found": sum(1 for _, t in res["deck"]["titles"]
                                                             if any(t.split(":")[0][:20] in p.get_text() for p in d))}
    OUT.write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    n_issues = sum(len(s["issues"]) for s in res["slides"])
    print(f"slides {res['deck']['slide_count']}, issues {n_issues}, order ok {order_ok}, "
          f"limits ok {all(v[3] for v in limits.values())}, low-res {len(low)}, pdf pages {d.page_count}")
    for s in res["slides"]:
        for it in s["issues"]:
            print(f"  slide {s['n']}: {it}")
    print("fonts:", res["deck"]["fonts"])
    print("pdf fonts:", fonts[:12])


if __name__ == "__main__":
    main()
