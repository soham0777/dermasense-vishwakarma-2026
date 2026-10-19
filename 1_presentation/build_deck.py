"""
Build the FINAL Stage 2 deck on the official Vishwakarma template (Template (2), never edited in place).

Template rule followed: sections in the template's order and within its slide limits; every content slide reuses the
template's own example-slide design for that section (background, logos, title position, panels, icons, rationale
bar). The guideline-only instruction slides are dropped, as the template intends; their requirements are mapped
slide by slide in 1_presentation/PPT_Asset_Index.md.

  1 Application ID (cover)            2 Problem Context                3-4 User Journey (before / after)
  5-6 System Architecture & Specs      7-8 CAD Design & Mechanical      9-10 Electronics Schematics & BOM
  11-12 Maker Journey                  13 Safety & Hazard Mitigation    14 Roadmap
  15 Technical Challenges & Risks      16 Potential Solutions & Risk Mitigation
  17 References & Sources              18 Image Sources

Evidence discipline: every slide carries one status pill (green = completed & verified in software or by source,
amber = designed / physical validation pending, grey = future scope, blue = research or documented decision).
Numbers come from project files or cited sources; targets are called targets.
Run:  python 1_presentation/render_cad.py  (SolidWorks)  ->  make_figures.py  ->  this script
"""
import csv
import re
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = HERE / "template" / "Vishwakarma Awards 2026-27 Stage 2 Submission Template (2).pptx"
OUTDIR = ROOT / "FINAL_SUBMISSION"                     # the only place the deck lives
OUT = OUTDIR / "Vishwakarma_Stage2_Final.pptx"
FIG = HERE / "figures"

NAVY, BODY, MUTED, GREEN, BLUE = "0F172A", "334155", "64748B", "059669", "0369A1"
PILL = {"green": ("047857", "D1FAE5"), "amber": ("B45309", "FEF3C7"), "grey": ("475569", "E2E8F0"), "blue": ("0369A1", "E0F2FE")}
TITLE_FONT, BODY_FONT = "Poppins SemiBold", "DM Sans"
SUP = dict(zip(range(1, 16), "¹ ² ³ ⁴ ⁵ ⁶ ⁷ ⁸ ⁹ ¹⁰ ¹¹ ¹² ¹³ ¹⁴ ¹⁵".split()))
NOTES = {}


# ================================================================== low-level helpers
def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def delete(slide, *sids):
    for sid in sids:
        el = by_id(slide, sid)._element
        el.getparent().remove(el)


def move(slide, sid, x=None, y=None, w=None, h=None):
    sh = by_id(slide, sid)
    if x is not None:
        sh.left = Inches(x)
    if y is not None:
        sh.top = Inches(y)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)
    return sh


def runs(p, parts, size, color=BODY, font=BODY_FONT):
    """parts: list of (text, bold[, color]) appended as runs to paragraph p."""
    for part in parts:
        text, bold = part[0], part[1]
        col = part[2] if len(part) > 2 else color
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = font
        f.color.rgb = RGBColor.from_string(col)


def textbox(slide, x, y, w, h, paras, size=10.0, color=BODY, font=BODY_FONT, align=None, space=4, anchor=MSO_ANCHOR.TOP,
            bullets=False, line=1.08):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, bold[, color]) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        parts = [(para, False)] if isinstance(para, str) else list(para)
        if bullets:
            parts = [("▪  ", False, GREEN)] + parts
        runs(p, parts, size, color, font)
        p.space_after = Pt(space)
        p.line_spacing = line
        if align:
            p.alignment = align
    return tb


def set_text(shape, paras, size, color=BODY, font=BODY_FONT, align=None):
    tf = shape.text_frame
    body = tf._txBody
    for p in body.findall(qn("a:p"))[1:]:
        body.remove(p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        parts = [(para, False)] if isinstance(para, str) else list(para)
        runs(p, parts, size, color, font)
        if align:
            p.alignment = align
    return shape


def title(slide, sid, text):
    sh = by_id(slide, sid)
    set_text(sh, [[(text, False)]], 20.25, NAVY, TITLE_FONT)
    sh.width = Inches(9.4)


def pill(slide, text, kind):
    fg, bg = PILL[kind]
    w = 0.30 + 0.068 * len(text)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.75 - w), Inches(0.50), Inches(w), Inches(0.27))
    shp.adjustments[0] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(bg)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, [(text, True)], 7.5, fg, BODY_FONT)
    return shp


def picture_in(slide, path, x, y, w, h, frame=None):
    """Add an image scaled to fit (x, y, w, h) inches, centred; optionally resize a template frame around it."""
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    if frame is not None:
        move(slide, frame, x=x - 0.01, y=y - 0.01, w=w + 0.02, h=h + 0.02)
    return pic


def _border(cell, color="E2E8F0", width=0.75, sides=("B",)):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("L", "R", "T", "B"):
        tag = qn(f"a:ln{side}")
        for old in tcPr.findall(tag):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, tag, w=str(int(width * 12700)))
        if side in sides:
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            etree.SubElement(ln, qn("a:noFill"))


def table(slide, x, y, w, rows, widths, size=8.6, head_size=8.2, row_h=0.26, bold_first_col=True, colors=None):
    """Native table in the template's flat style: bold header, hairline separators, no fills."""
    n, m = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    t = gf.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"       # "No Style, No Grid"
    total = sum(widths)
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(w * cw / total)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.background()
            c.margin_left = c.margin_right = Inches(0.04)
            c.margin_top = Inches(0.025)
            c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            head = i == 0
            col = NAVY if head or (j == 0 and bold_first_col) else BODY
            if colors and (i, j) in colors:
                col = colors[(i, j)]
            if isinstance(val, str):
                parts = [(val, head or (j == 0 and bold_first_col), col)]
            elif isinstance(val, tuple):                      # one styled cell: (text, bold, colour)
                parts = [val]
            else:
                parts = val
            runs(p, parts, head_size if head else size)
            _border(c, "CBD5E1" if head else "E2E8F0", 1.0 if head else 0.6)
    return gf


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def duplicate(prs, src):
    """Copy a slide (shapes, images, background) to the end of the deck."""
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        rid_map[rid] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else new.part.relate_to(rel.target_part, rel.reltype))
    for el in src.shapes._spTree:
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            new.shapes._spTree.append(deepcopy(el))
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for el in new.shapes._spTree.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(r_ns) and val in rid_map:
                el.set(attr, rid_map[val])
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, deepcopy(bg))
    return new


def reorder_and_prune(prs, keep):
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(s.rId): s for s in lst}
    wanted = [by_part[s.part] for s in keep]
    for s in list(lst):
        if s not in wanted:
            prs.part.drop_rel(s.rId)
        lst.remove(s)
    for s in wanted:
        lst.append(s)


def bom_groups():
    rows = {r["id"]: r for r in csv.DictReader(open(ROOT / "6_electronics" / "bom_v0_1.csv", encoding="utf-8"))}

    def rng(ids):
        lo = hi = 0.0
        for i in ids:
            a, b = (rows[i]["est_unit_cost_inr"].split("-") + [None])[:2]
            a = float(a.split()[0])
            b = float(b) if b else a
            q = float(rows[i]["qty"])
            lo, hi = lo + a * q, hi + b * q
        return lo, hi
    groups = [("ESP32-C3 SuperMini", "USB-C, native USB serial", "1", "Probots · quoted", ["R01"]),
              ("ADS1115 16-bit ADC", "I²C, auto-range", "1", "Probots · quoted", ["R02"]),
              ("P75-B1 pogo pins", "1 mm, pack of 100", "100", "Hubtronics · quoted", ["M01"]),
              ("Resistors, C0G cap, LEDs, board", "0.1 % + 1 % resistors, headers", "set", "local / online",
               ["R03", "R04", "R05", "R06", "R07"]),
              ("Clamp + bench rig", "PETG print · acrylic · M3/M4 · dowels", "set", "lab / local", ["M02", "M03", "M04", "M05", "P04"]),
              ("Strip materials", "polyimide film · seal tape · silver paint", "set", "lab / online", ["P01", "S01", "P02"]),
              ("Test consumables", "1 % resistor set · IPA · gloves · saline", "set", "local", ["P03", "P05", "P06", "P07"])]
    out, tl, th = [], 0.0, 0.0
    for name, spec, qty, src, ids in groups:
        lo, hi = rng(ids)
        tl, th = tl + lo, th + hi
        out.append((name, spec, qty, src, f"{lo:,.0f}" if lo == hi else f"{lo:,.0f}–{hi:,.0f}"))
    return out, (tl, th)


# ================================================================== slide content
# Design rule (declutter, 16 Oct): one message per slide, one large visual, at most 4 short points, one-line takeaway in
# the template's rationale bar. Details live in the speaker notes and in PPT_Claim_and_Reference_Check.md.
B_SIZE = 12.0          # panel bullet size (was ~9.5)
BAR_SIZE = 9.6


def pts(*items):
    """Short bullet points: (label, text) pairs."""
    return [[(f"{a}  ", True, NAVY), (b, False)] for a, b in items]


def cover(sl):
    by_id(sl, 86).top, by_id(sl, 86).height = Inches(4.30), Inches(0.62)
    set_text(by_id(sl, 86), [[("VK26-IITI-HT-0050", False)]], 24, BLUE, TITLE_FONT, PP_ALIGN.CENTER)
    move(sl, 86, x=2.2, w=8.93)
    textbox(sl, 1.4, 5.02, 10.53, 0.5, [[("Dermasense: a laser-written strip that reads a dressing's own 'change line'",
                                          False)]], 16, NAVY, TITLE_FONT, PP_ALIGN.CENTER, space=2)
    textbox(sl, 1.4, 5.62, 10.53, 0.8, [
        [("Soham Prabhakar Kadu · Bhakti Ajay Kadam", False)],
        [("Health Tech: Emergency Response & Remote Monitoring", False)]], 11, BODY, align=PP_ALIGN.CENTER, space=3)
    textbox(sl, 1.4, 6.45, 10.53, 0.3, [[("Evidence key:  ", True, MUTED), ("● verified", True, GREEN),
                                         ("   ● designed, lab test pending", True, "B45309"),
                                         ("   ● future scope", True, "475569")]], 8.5, MUTED, align=PP_ALIGN.CENTER)
    notes(sl, "Core innovation: two laser-carbonised electrode zones, at the pad centre and on the dressing maker's change "
              "line, let a clip-on reader say OK, CHANGE or CHECK without lifting the dressing.")


def e4_slide(sl, ttl, pill_txt, pill_kind, metrics, fig, caption, heading, bullets, bar, note):
    title(sl, 133, ttl)
    delete(sl, 134, 152, 153, 154, 155, 148, 149, 150, 151, 145)
    pill(sl, pill_txt, pill_kind)
    for (vid, lid), (val, lab) in zip(((136, 137), (138, 139), (140, 141), (142, 143)), metrics):
        v = set_text(by_id(sl, vid), [[(val, False)]], 15 if len(val) < 12 else 13, BLUE, TITLE_FONT, PP_ALIGN.CENTER)
        v.top, v.height = Inches(1.12), Inches(0.33)
        lb = set_text(by_id(sl, lid), [[(lab, False)]], 7.6, MUTED, "DM Sans Medium", PP_ALIGN.CENTER)
        lb.top, lb.height = Inches(1.47), Inches(0.24)
    picture_in(sl, FIG / fig, 0.59, 1.91, 5.92, 3.93, frame=132)
    cap = by_id(sl, 144)
    set_text(cap, [[(caption, False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.90), Inches(0.22)
    move(sl, 131, y=1.90, h=4.25)
    set_text(by_id(sl, 147), [[(heading, True)]], 11.5, NAVY)
    by_id(sl, 147).width = Inches(5.3)
    textbox(sl, 7.07, 2.62, 5.5, 3.4, bullets, B_SIZE, bullets=True, space=12)
    set_text(by_id(sl, 135), bar, BAR_SIZE, BODY)
    move(sl, 135, w=11.7, h=0.36)
    notes(sl, note)


def problem(sl):
    e4_slide(sl, "Problem Context", "PUBLISHED EVIDENCE · CITED", "blue",
             [("44.9 %", f"OF 588 CHANGES MADE AT OPTIMAL MOISTURE{SUP[1]}"),
              ("≈ 2 cm", f"MAKER'S 'CHANGE NOW' LINE FROM PAD EDGE{SUP[2]}"),
              ("1 point", f"WHERE TODAY'S MOISTURE SENSOR READS{SUP[1]}"),
              ("Uneven", f"IN-VITRO MOISTURE MAP, 8 POINTS{SUP[3]}")],
             "fig_problem.png", "Figure: the change line is invisible under a cover (to scale, our design).",
             "Target Beneficiary & Problem Statement",
             pts(("Who:", "nurses and home carers of chronic wounds under foam dressings."),
                 ("Pain:", f"the maker's 'change now' sign is visual{SUP[14]}, and the dressing is covered."),
                 ("Today:", "change by the clock, by lifting the cover, or after a leak."),
                 ("Gap:", f"in our search, no low-cost sensor reads that change line{SUP[4]}{SUP[5]}.")),
             [[("Problem: ", True, NAVY), ("carers cannot see when a covered dressing is really full.", False)]],
             "Target user: nurses and carers (diabetic foot, venous leg, pressure ulcers). Public search: WoundSense reads one "
             "point [1]; a lab study mapped 8 points [3]; diaper wetness strips with clip-on readers exist [4]; smart bandages "
             "use custom electronics [5]. Visual change indicators exist (ALLEVYN LIFE: 'At 75%, it's time to change' [14]) but "
             "need the dressing face to be visible. Why now: some makers define 'full' as a distance from the pad edge [2] and conductive "
             "carbon can be laser-written on polyimide in one pass [7]. Project: Dermasense, VK26-IITI-HT-0050, Health Tech.")


def journey_before(sl):
    e4_slide(sl, "User Journey: Before State", "CURRENT STATE · CITED", "blue",
             [("44.9 %", f"CHANGES WHILE MOISTURE WAS STILL OPTIMAL{SUP[1]}"),
              ("≈ 18 min", f"A DRESSING VISIT, PATIENT-REPORTED{SUP[13]}"),
              ("₹240–510", f"ONE 10 × 10 CM FOAM DRESSING, RETAIL{SUP[15]}"),
              ("0.1–1.3 g/cm²·day", f"REPORTED EXUDATE RATES, 2 METHODS{SUP[6]}")],
             "fig_journey_before.png", "Figure: today's pain-point map (cited numbers only).",
             "Current Pain Point Mapping",
             pts(("Time:", f"every check means unwrap, look, re-secure: about 18 min a visit{SUP[13]}."),
                 ("Waste:", f"44.9 % of changes came while moisture was still optimal{SUP[1]}; each discards a "
                            f"dressing{SUP[15]}."),
                 ("Risk:", "a late change can leak: wet skin, soiled clothing."),
                 ("Cause:", f"reported exudate rates span over 10×{SUP[6]}; a clock cannot follow.")),
             [[("Quantified Inefficiency: ", True, NAVY), ("published figures only (a UK patient survey, Indian retail "
                                                           "prices); we will time real checks with nurses.", False)]],
             "Sources: Milne 2016 [1]; Dealey 2006 [6] (0.10 and 0.21 g/cm²/24 h by one method, mean 1.3 by the other); "
             "UK survey of >200 patients commissioned by Mölnlycke, 2018 [13]: average dressing time 18 min, about 5 changes a "
             "week (patient-reported, industry-sponsored); Truemeds listings [15]: Mepilex Lite 10x10 cm Rs 242, Mepilex Border "
             "10x10 cm Rs 514 (pack size unclear). Not yet measured: nurse minutes per check in Indian clinics.")


def journey_after(sl):
    e4_slide(sl, "User Journey: After State (Intended)", "INTENDED · NOT CLINICALLY VALIDATED", "amber",
             [("up to 44.9 %", f"EARLY CHANGES A RELIABLE ALERT COULD AVOID{SUP[1]}"),
              ("≤ ₹100", "TARGET STRIP COST (ESTIMATE)"),
              ("±10 %", "TARGET ALERT ACCURACY (BENCH, E4)"),
              ("0", "CONTACT WITH THE WOUND BED (BY DESIGN)")],
             "fig_journey_after.png", "Figure: the intended workflow; bench-testable, not clinically validated.",
             "Future Solution Mapping & Quantified Efficiency",
             pts(("Same routine:", "strip on the pad, reader clips on."),
                 ("Clear signal:", "CHANGE when the line is wet; CHECK if it wets first."),
                 ("Remote:", "the state is broadcast over Bluetooth to a phone."),
                 ("Payback:", f"a ≤ ₹100 strip pays for itself if 1 in 2–5 changes is avoided{SUP[15]}.")),
             [[("Design Principle: ", True, NAVY), ("only film and carbon are disposable; the strip never touches the "
                                                    "wound.", False)]],
             "Targets and projections, not results. Payback: strip <= Rs 100 (estimate) / dressing Rs 242-514 [15] = 0.19-0.41, "
             "i.e. one avoided change in every 2.4-5.1 dressings. 'Up to 44.9 %' assumes every change waits for CHANGE; some "
             "changes are made for other clinical reasons, so the real figure is lower and must be measured. Bluetooth: compiled, "
             "not yet tested on hardware (bench test T11). FAULT covers a broken strip or reader error.")


def arch(sl):
    title(sl, 186, "System Architecture & Specs")
    delete(sl, 187, 190, 191, *range(194, 212))
    pill(sl, "DESIGNED · LAB TEST PENDING", "amber")
    picture_in(sl, FIG / "fig_block.png", 0.59, 1.06, 5.92, 4.24, frame=185)
    cap = by_id(sl, 189)
    set_text(cap, [[("Figure: subsystems; ✓ = verified in software, ◐ = hardware still to prove.", False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.34), Inches(0.24)
    set_text(by_id(sl, 193), [[("Target Specifications & Hardware", True)]], 11.5, NAVY)
    table(sl, 7.05, 1.70, 5.48, [
        ("Parameter", "Design value"),
        ("Sensing zones", "2: pad centre C, change line A"),
        ("Wet / dry", "< 100 kΩ / > 300 kΩ, 3 scans"),
        ("Drive", "bipolar, ≤ 33 µA, no net DC"),
        ("ADC / MCU", "ADS1115 16-bit (I²C) · ESP32-C3"),
        ("Contacts", "4 pogo pins + 2 dowels"),
        ("Output", "LEDs · USB log · BLE broadcast"),
    ], [1.6, 3.2], size=10.5, head_size=9.5, row_h=0.5)
    set_text(by_id(sl, 188), [[("Architecture Rationale: ", True, NAVY),
                               ("only film and carbon are disposable; every electronic part is reused.", False)]], BAR_SIZE, BODY)
    move(sl, 188, w=11.7)
    notes(sl, "Electrode lines/gaps 1.0 mm (provisional until E2). Windows 6 x 3.9 mm. Expected wet 4.2-14.8 kOhm, dry >= 1 MOhm "
              "(calc_readout.py). Self-check: 100 kOhm 0.1 % resistor (+/-5 %) and COM loop -> FAULT. BLE: one 13-byte "
              "advertisement per scan (ble_beacon.h), compiled, not run on hardware. Sources: config.h, "
              "strip_v0_1_geometry.json.")


def e8_slide(sl, ttl, pill_txt, pill_kind, fig, caption, heading, bullets, bar, note):
    title(sl, 238, ttl)
    delete(sl, 239, 242, 245, 246, 247, 248, 249, 250, 251, 252, 253, 254)
    pill(sl, pill_txt, pill_kind)
    picture_in(sl, FIG / fig, 0.59, 1.06, 5.92, 4.24, frame=237)
    cap = by_id(sl, 241)
    set_text(cap, [[(caption, False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.36), Inches(0.22)
    move(sl, 236, h=4.85)
    set_text(by_id(sl, 244), [[(heading, True)]], 11.5, NAVY)
    by_id(sl, 244).width = Inches(5.3)
    textbox(sl, 7.07, 1.72, 5.5, 4.0, bullets, B_SIZE, bullets=True, space=14)
    set_text(by_id(sl, 240), bar, BAR_SIZE, BODY)
    move(sl, 240, w=11.7, h=0.36)
    notes(sl, note)


def firmware(sl):
    e8_slide(sl, "System Architecture & Specs: Firmware & Data Flow", "VERIFIED IN SOFTWARE · HARDWARE PENDING", "green",
             "fig_dataflow.png", "Figure: one reader scan, from measurement to state.",
             "Firmware Logic & Data Flow",
             pts(("Measure:", "each zone driven both ways, so no net DC."),
                 ("Read:", "16-bit ADC, scaled to the 3V3 rail."),
                 ("Decide:", "wet < 100 kΩ, dry > 300 kΩ, 3 scans in a row."),
                 ("Self-check:", "a known 100 kΩ every scan, else FAULT."),
                 ("Share:", "CSV over USB, and a Bluetooth broadcast to a phone.")),
             [[("Software vs Hardware: ", True, NAVY), ("23 compile-time and 16 PC tests pass; not yet run on a wired board.",
                                                        False)]],
             "ADS1115 single-shot 128 SPS, auto-range; A1 reads the 3V3 rail. Rx from both phases, averaged, scaled by the check "
             "resistor (+/-5 % else FAULT), minus 2 kOhm protection. Idle pins high-Z; zones frozen during FAULT. Compiles for "
             "ESP32-C3 with 0 warnings (650,349 bytes, 49 % of flash, with Bluetooth). BLE payload: state, wet flags, scan number, "
             "both zone resistances; unencrypted, test company ID 0xFFFF, bench only. Decoder: 7_firmware/tools/decode_beacon.py.")


def cad(sl):
    e8_slide(sl, "CAD Design & Mechanical Specs", "DESIGNED · CAD VERIFIED · NOT YET BUILT", "amber",
             "fig_cad.png", "Figure: our SolidWorks 2019 models: exploded stack, clamp, drawing, assembly.",
             "Materials & Mechanical Checks",
             pts(("Materials:", "polyimide strip · PETG clamp · acrylic rig."),
                 ("Contact:", "4 pogo pins, 0.72 N each (calculated)."),
                 ("Alignment:", "2.8× tolerance margin on 3 mm pads."),
                 ("Verified:", "99 automated CAD checks pass, 0 fail."),
                 ("Provisional:", "strip sizes wait for E2 laser data.")),
             [[("Design Iteration: ", True, NAVY), ("a clash check moved the seal tape off the clamp (decision D-013).",
                                                    False)]],
             "Hand calculations (calc_mechanical.py): 1.48 mm pin compression, 2.9 N total, lid deflection 11 um, tolerance stack "
             "0.49 mm vs +/-1.35 mm, 0.62 % bending strain at 10 mm radius, 20 MPa vs 231 MPa at the dowels. No FEA: loads < 3 N. "
             "9 parts, 1 assembly, 7 A3 drawings, STEP. Seal clash was 9.0 mm3; now 0.")


def strip(sl):
    e8_slide(sl, "CAD Design & Mechanical Specs: Sensing Strip & Laser Process", "DESIGNED · LAB TEST PENDING", "amber",
             "fig_strip.png", "Figure: our laser files, drawn from the DXFs. Not yet lased.",
             "Fabrication & Process Validation (E2)",
             pts(("Material:", f"laser-carbonised, LIG-like layer on polyimide{SUP[7]}."),
                 ("Reference:", "setting '12', most likely % PWM, not watts."),
                 ("E2 plan:", "7 coupons, 2 days, 72 pre-registered samples."),
                 ("Pass rules:", "CV ≤ 20 %, gaps > 3 MΩ, fixed beforehand."),
                 ("Sealed:", f"leads under tape; LIG made in air wets easily{SUP[9]}.")),
             [[("Laser Files Verified: ", True, NAVY), ("35/35 file checks pass; LaserCAD import is checked at the machine.",
                                                        False)]],
             "Graphene not confirmed until Raman. Machine (LaserCAD config): 1200 x 900 mm bed, 20 kHz PWM, 98 % cap. Logged per "
             "layer: min/max power, speed, 0.1 mm interval, passes, scan direction, focus, tube current. E2 = 4 bracket coupons "
             "(power 10/12/14, speed x0.67/1/1.5) + 3 resolution coupons.")


def schematic(sl):
    title(sl, 281, "Electronics Schematics & BOM")
    delete(sl, 282, 285, 286, *range(289, 324))
    pill(sl, "ERC VERIFIED · NOT YET WIRED", "green")
    picture_in(sl, FIG / "fig_schematic.png", 0.59, 1.06, 5.92, 4.24, frame=280)
    cap = by_id(sl, 284)
    set_text(cap, [[("Figure: our KiCad 10 schematic of the reader (exported by kicad-cli).", False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.36), Inches(0.24)
    set_text(by_id(sl, 288), [[("Key Connections", True)]], 11.5, NAVY)
    textbox(sl, 7.07, 1.72, 5.5, 4.0, pts(("Drive:", "GPIO3 / GPIO4 → 1 kΩ → zones C, A."),
                                         ("Sense:", "node N → 100 kΩ reference → ADC A0."),
                                         ("Check:", "100 kΩ on GPIO20; strip loop on GPIO5."),
                                         ("Strip:", "4 pogo pins on 3 × 6 mm pads."),
                                         ("Power:", "USB-C → 3V3 regulator on the module.")),
            B_SIZE, bullets=True, space=14)
    set_text(by_id(sl, 283), [[("Verification: ", True, NAVY),
                               ("ERC 0 violations; netlist matches firmware and strip pads (0 mismatches).", False)]],
             BAR_SIZE, BODY)
    move(sl, 283, w=11.7)
    notes(sl, "Full pin map: 7_firmware/WIRING.md (generated from the netlist). I2C on GPIO6/7 (ADS1115 0x48); A1 = 3V3 rail; "
              "LEDs on GPIO0/1/21; strapping pins GPIO2/8/9 unused.")


def bom(sl):
    title(sl, 281, "Electronics Schematics & BOM: Bill of Materials")
    delete(sl, 282, 285, 286, *range(289, 324))
    pill(sl, "3 PRICES QUOTED · NOTHING ORDERED", "amber")
    picture_in(sl, FIG / "fig_window.png", 0.59, 1.06, 5.92, 4.24, frame=280)
    cap = by_id(sl, 284)
    set_text(cap, [[("Figure: why a 100 kΩ reference (calculated ranges).", False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.36), Inches(0.24)
    groups, (tl, th) = bom_groups()
    rows = [("Item", "Qty", "Source", "₹")] + [(g[0], g[2], g[3], g[4]) for g in groups] + \
           [(("Total", True, NAVY), "", "", (f"≈ {tl:,.0f}–{th:,.0f}", True, NAVY))]
    set_text(by_id(sl, 288), [[("Prototype Bill of Materials", True)]], 11.5, NAVY)
    table(sl, 7.05, 1.66, 5.48, rows, [2.3, 0.45, 1.45, 1.28], size=9.8, head_size=9.2, row_h=0.41)
    set_text(by_id(sl, 283), [[("Procurement Status: ", True, NAVY),
                               ("reader parts and pins quoted 19 Oct, in stock in India; the rest are estimates.", False)]], BAR_SIZE, BODY)
    move(sl, 283, w=11.7)
    notes(sl, "Line items: 6_electronics/bom_v0_1.csv (23 lines, each traced to the file that needs it). First order: silver "
              "paint, 1 % resistor set, 0.1 % 100 kOhm resistors, ADS1115, ESP32-C3. Quoted 19 Oct 2026 incl. GST: Probots PRO4607 "
              "ESP32-C3 SuperMini Rs 242; Probots PRO789 ADS1115 Rs 125; Hubtronics 05880 P75-B1 pack of 100 Rs 354. Total leaves "
              "out the borrowed meter, USB cable and E4 dressings. Ranges = sum of min/max estimates.")


def e12_slide(sl, ttl, pill_txt, pill_kind, fig, caption, heading, bullets, bar, note):
    title(sl, 366, ttl)
    delete(sl, 367, 370, 373, 374, 375, 376, 377, 378, 379, 380)
    pill(sl, pill_txt, pill_kind)
    picture_in(sl, FIG / fig, 0.59, 1.06, 5.92, 4.24, frame=365)
    cap = by_id(sl, 369)
    set_text(cap, [[(caption, False)]], 7.6, MUTED)
    cap.top, cap.height = Inches(5.36), Inches(0.22)
    move(sl, 364, h=4.95)
    set_text(by_id(sl, 372), [[(heading, True)]], 11.5, NAVY)
    by_id(sl, 372).width = Inches(5.3)
    textbox(sl, 7.07, 1.72, 5.5, 4.2, bullets, B_SIZE, bullets=True, space=14)
    set_text(by_id(sl, 368), bar, BAR_SIZE, BODY)
    move(sl, 368, w=11.7, h=0.36)
    notes(sl, note)


def maker(sl):
    e12_slide(sl, "Maker Journey", "GENUINE ARTIFACTS · NO STAGED EVIDENCE", "green", "fig_maker.png",
              "Figure: our SolidWorks window and verbatim tool output (14–15 Oct 2026).",
              "Authenticity: What Each Artifact Proves",
              pts(("CAD:", "our assembly open in SolidWorks, 13 parts."),
                  ("Checks:", "CAD, ERC, compile and tests all re-runnable."),
                  ("History:", "every step a git commit since 14 Oct."),
                  ("Photos:", "none yet; the first laser session is next.")),
              [[("Raw Repository Linkage: ", True, NAVY), ("[team to insert the GitHub / Drive link before upload]", False)]],
              "Native files in the repository: .SLDPRT, .SLDASM, .SLDDRW, .STEP, .kicad_sch, .ino, .dxf. CAD re-check: 99 pass, "
              "0 fail. KiCad ERC log: 0 violations. Faculty sign-off follows the E2 review. We do not stage or generate lab "
              "photos.")


def iterations(sl):
    e12_slide(sl, "Maker Journey: Engineering Iterations", "DOCUMENTED DECISIONS", "blue", "fig_decision.png",
              "Figure: concept scores and the iteration path (each a dated commit).",
              "Stage 1 to Stage 2 Iterations",
              pts(("Pivot:", f"therapy bandage (already shown in labs{SUP[5]}) → one job: is it full?"),
                  ("Choice:", "two-zone strip (4.36) beat one zone (4.29): it spots leaks."),
                  ("Kept:", "Stage 1's Bluetooth telemetry, now a state broadcast."),
                  ("Corrected:", f"Stage 1 called Kapton 'medical-grade'; DuPont cautions otherwise{SUP[11]}."),
                  ("Redesign:", "FMEA risk → v0.2 loops each sense lead.")),
              [[("Prior Art: ", True, NAVY), (f"diaper strips{SUP[4]} and visual change indicators{SUP[14]} exist; our "
                                              "novelty is reading the change line electrically, with an order check.", False)]],
              "5 concepts scored on 13 weighted criteria (3_decisions/decision_matrix.md). Verification fixes: seal moved off the "
              "clamp; STEP was in cm, now mm; drawings rebuilt first-angle. FMEA F04 (cracked sense lead reads dry, RPN 224). "
              "4 zones were cut to 2 until E2 proves finer lines. Stage 1 also estimated '60 % fewer changes'; we dropped that "
              "number (no source). IP: concept and files are the team's own; no patent filed yet.")


def safety(sl):
    title(sl, 421, "Safety & Hazard Mitigation")
    delete(sl, 422, 423, *range(426, 435), 437, 438, 441, 442, 443, 444, 445, 446)
    pill(sl, "RISKS ESTIMATED · TESTS PLANNED", "amber")
    move(sl, 418, h=2.30)
    set_text(by_id(sl, 425), [[("Design FMEA: top risks of 15 (ratings estimated before tests)", True)]], 11.5, NAVY)
    by_id(sl, 425).width = Inches(7.4)
    table(sl, 0.82, 1.62, 8.15, [
        ("Failure mode", "RPN", "Fix"),
        ("Cracked sense lead reads 'dry'", ("224", True, "DC2626"), "v0.2: loop each lead to a 2nd pad"),
        ("Front path differs from model", ("210", True, "DC2626"), "test on real foam dressings (E4)"),
        ("Humid dressing reads 'wet'", ("150", True, "DC2626"), "measure dry baselines (E4)"),
    ], [3.0, 0.6, 3.4], size=10.5, head_size=9.5, row_h=0.36)
    picture_in(sl, FIG / "fig_fmea.png", 9.15, 1.25, 3.45, 2.06)
    for pid, iid, hid in ((419, 435, 436), (420, 439, 440)):
        move(sl, pid, y=3.50, h=2.95)
        move(sl, iid, y=3.74)
        move(sl, hid, y=3.71, w=5.3)
    set_text(by_id(sl, 436), [[("Fail-safes", True)]], 11.5, NAVY)
    set_text(by_id(sl, 440), [[("Standards (considered, not claimed)", True)]], 11.5, NAVY)
    textbox(sl, 0.84, 4.10, 5.5, 2.3, pts(("Bench only:", "USB reader, never on a person."),
                                         ("No net DC:", "bipolar drive, idle pins off."),
                                         ("Self-check:", "every scan, else FAULT."),
                                         ("No wound contact:", "strip on the pad's outer face.")),
            11, bullets=True, space=7)
    textbox(sl, 7.07, 4.10, 5.5, 2.3, pts(("ISO 14971", "risk management"), ("IEC 60601-1", "future on-body reader"),
                                         ("ISO 10993-1", f"before any skin contact{SUP[11]}"),
                                         ("EN 13726", f"simulated exudate for tests{SUP[10]}"),
                                         ("IEC 60825-1", "laser work during fabrication")),
            11, bullets=True, space=5)
    notes(sl, "Full FMEA: 3_decisions/fmea_v0_1.md (15 modes). Laser safety: signed checklist each session "
              "(4_sensor_strip/e2_experiment/E2_safety_checklist.md). Bench test T8 checks average DC < 5 mV.")


def roadmap(sl):
    title(sl, 421, "Roadmap")
    delete(sl, 422, 423, *range(426, 435), 437, 438, 441, 442, 443, 444, 445, 446)
    pill(sl, "FUTURE SCOPE · GATED PLAN", "grey")
    move(sl, 418, h=2.62)
    set_text(by_id(sl, 425), [[("4-Month Plan on the programme's own dates: each gate must pass first", True)]], 11.5, NAVY)
    by_id(sl, 425).width = Inches(9.0)
    picture_in(sl, FIG / "fig_roadmap.png", 0.72, 1.55, 11.9, 2.05)
    for pid, iid, hid in ((419, 435, 436), (420, 439, 440)):
        move(sl, pid, y=3.82, h=2.62)
        move(sl, iid, y=4.06)
        move(sl, hid, y=4.03, w=5.3)
    set_text(by_id(sl, 436), [[("Build & Bench (Months 1 - 2)", True)]], 11.5, NAVY)
    set_text(by_id(sl, 440), [[("Validate & Demo (Months 3 - 4)", True)]], 11.5, NAVY)
    textbox(sl, 0.84, 4.42, 5.5, 1.9, pts(("Nov:", "E2 laser tests (G1); build and bench-test the reader (G2)."),
                                         ("Early Dec:", "first strips on a dressing; 3-min prototype video.")),
            11.5, bullets=True, space=10)
    textbox(sl, 7.07, 4.42, 5.5, 1.9, pts(("Dec:", "strip v0.2 + printed clamp; E4 wet tests (G3)."),
                                         ("Jan:", "repeatability; virtual semi-final with a live demo."),
                                         ("Stage 3:", "OK → CHANGE live on a dressing at IIT Indore.")),
            11.5, bullets=True, space=10)
    notes(sl, "Gates: G1 E2 Rs CV <= 20 %; G2 reader bench test T0-T11; G3 E4 alert within +/-10 %; G4 live demo with logged "
              "raw data. Programme dates (vishwakarma-awards.org and the Internshala listing, read 19 Oct 2026): prototype video "
              "+ updated PPT in December (listing: 4 Dec 2026); Health Tech semi-final interviews in January; IIT Indore bootcamp "
              "and finale 26 Jan - 1 Feb 2027. No patient use. Raman/SEM via a partner lab if access is granted.")


def e18_cards(sl, ttl, pill_txt, pill_kind, cards, bar, note, card_h=3.05):
    title(sl, 482, ttl)
    delete(sl, 483, 485, 486, 487, 488, 489, 490, 491)
    pill(sl, pill_txt, pill_kind)
    for pid in (479, 480, 481):
        move(sl, pid, h=card_h)
    for (hid, x), (head, paras) in zip(((493, 0.82), (495, 4.95), (497, 9.09)), cards):
        set_text(by_id(sl, hid), [[(head, True)]], 11, NAVY)
        by_id(sl, hid).width = Inches(3.35)
        textbox(sl, x, 1.66, 3.45, card_h - 0.66, paras, 11, space=9)
    set_text(by_id(sl, 484), bar, BAR_SIZE, BODY)
    move(sl, 484, w=11.7, h=0.36)
    notes(sl, note)


def card(*items):
    return [[(f"{a} ", True, NAVY), (b, False)] for a, b in items]


def challenges(sl):
    e18_cards(sl, "Technical Challenges & Risks", "FOUND IN OUR OWN VERIFICATION", "blue", [
        ("Geometry: Seal Under Clamp", card(("Found:", "seal tape overlapped the clamp (9.0 mm³)."),
                                            ("Impact:", "would lift the strip off the pins."),
                                            ("How:", "automated clash check."))),
        ("Circuit: Hidden Lead Break", card(("Found:", "a cracked sense lead reads 'dry'."),
                                            ("Impact:", "CHANGE never shown, no alarm."),
                                            ("How:", "design FMEA (RPN 224)."))),
        ("Process: Unknown Laser Recipe", card(("Found:", f"'12' is likely a % command; tubes are unsteady at low power{SUP[8]}."),
                                               ("Impact:", "every strip size is provisional."),
                                               ("How:", "LaserCAD config audit."))),
    ], [[("Also Caught: ", True, NAVY), ("STEP exported in cm (now mm); a drawing showed 0.13 for 0.125 mm (fixed).", False)]],
        "Evidence: 5_cad/verify_cad_report.md (interference), 3_decisions/fmea_v0_1.md, 4_sensor_strip/machine_record.md.",
        card_h=2.75)
    picture_in(sl, FIG / "fig_challenges.png", 0.62, 4.05, 12.1, 2.0)


def solutions(sl):
    e18_cards(sl, "Potential Solutions & Risk Mitigation", "FIXED · DESIGNED · FUTURE", "grey", [
        ("Fixed & Verified", card(("Seal:", "starts at x = 23 mm; 0 clash."),
                                  ("Units:", "all STEP files in mm."),
                                  ("Checks:", "99 CAD checks pass."))),
        ("Designed, to Prove", card(("v0.2 strip:", "every sense lead looped."),
                                    ("Contact:", "silver pads, rounded pins."),
                                    ("Self-test:", "check resistor → FAULT."))),
        ("Fallbacks", card(("Wider lines:", "if E2 says so."),
                           ("Process:", "slower passes, log current."),
                           ("Concept:", "one-zone strip (score 4.29)."))),
    ], [[("Rule We Keep: ", True, NAVY), ("no performance, graphene or clinical claim until it is measured.", False)]],
        "Future scope details: 4_sensor_strip/e2_experiment/E2_measurements_and_unknowns.md.", card_h=2.45)
    box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(3.68), Inches(12.17), Inches(2.38))
    box.adjustments[0] = 0.06
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
    box.line.color.rgb = RGBColor.from_string("E2E8F0")
    box.shadow.inherit = False
    textbox(sl, 0.82, 3.80, 11.7, 0.3, [[("Future Scope & Validation Plan", True)]], 11.5, NAVY)
    textbox(sl, 0.82, 4.25, 5.75, 1.7, pts(("Experiments:", "E2 laser, reader bench test, E4 wet tests."),
                                          ("Material:", "Raman and SEM before any 'graphene' claim."),
                                          ("Reliability:", "bending, contact cycling, humidity drift.")),
            11, bullets=True, space=8)
    textbox(sl, 6.85, 4.25, 5.75, 1.7, pts(("Product:", "battery pod; phone app on the BLE broadcast; strips per dressing."),
                                          ("Before patients:", "ISO 10993, IEC 60601-1, ethics approval."),
                                          ("Not claimed:", "sensing performance, graphene, clinical benefit.")),
            11, bullets=True, space=8)


REFS = [
    "Milne S.D. et al., A wearable wound moisture sensor as an indicator for wound dressing change, Int. Wound J. 13(6):1309–1314 "
    "(2016). doi:10.1111/iwj.12521 · full text read",
    "Dressing instructions for use: some foam dressings say 'change when exudate approaches ~2 cm from the pad edge'; "
    "Mölnlycke Mepilex Border Flex brochure uses strike-through reaching 3 edges · manufacturer documents, read as excerpts",
    "McColl D., Cartlidge B., Connolly P., Real-time monitoring of moisture levels in wound dressings in vitro, Int. J. Surg. "
    "5(5):316–322 (2007). doi:10.1016/j.ijsu.2007.02.008 · abstract",
    "US 5,469,145 and US 5,838,240 (wet diaper detector: detachable strip + clip-on detector); US 9,291,589; US 10,928,344 · "
    "patent summaries, claims not read",
    "Jiang Y. et al., Wireless, closed-loop, smart bandage with integrated sensors and stimulators, Nat. Biotechnol. (2022). "
    "doi:10.1038/s41587-022-01528-3 · listing + coverage",
    "Dealey C., Cameron J., Arrowsmith M., A study comparing two objective methods of quantifying the production of wound "
    "exudate, J. Wound Care 15(4):149–153 (2006) · abstract (methods disagree)",
    "Lin J. et al., Laser-induced porous graphene films from commercial polymers, Nat. Commun. 5:5714 (2014). "
    "doi:10.1038/ncomms6714 · full text read",
    "CO₂ laser tube current vs PWM setting, engineering measurements, softsolder.com (2022) · not peer-reviewed",
    "Li Y. et al., Laser-induced graphene in controlled atmospheres: from superhydrophilic to superhydrophobic surfaces, "
    "Adv. Mater. 29:1700496 (2017). doi:10.1002/adma.201700496 · abstract",
    "EN 13726-1:2002 / EN 13726:2023, Test methods for wound dressings (absorbency; 'Solution A') · catalogue entry; "
    "online link not reachable from our network",
    "DuPont position statement on the use of DuPont materials in medical applications (Kapton caution) · web page",
    "Texas Instruments ADS1115 datasheet (SBAS444); Espressif ESP32-C3 datasheet · vendor documents",
    "Wounds UK news (22 Aug 2018) on a Mölnlycke-commissioned survey of >200 UK patients with chronic wounds: average "
    "dressing time 18 min · patient-reported, industry-sponsored; web page read",
    "Smith+Nephew, ALLEVYN LIFE product page: visual change indicator ('At 75%, it's time to change your dressing') · "
    "manufacturer web page, read 19 Oct 2026",
    "Truemeds.in retail listings: Mepilex Lite 10 × 10 cm ₹242; Mepilex Border 10 × 10 cm ₹514 (pack size unclear) · "
    "retail pages, read 19 Oct 2026",
]


def references(sl):
    delete(sl, 504, 505, 506, 507, 508, 509, 510, 511, 512, 516)
    set_text(by_id(sl, 513), [[("References & Sources", False)]], 20.25, NAVY, TITLE_FONT)
    pill(sl, "SOURCES WITH READING DEPTH", "blue")
    textbox(sl, 0.62, 1.12, 12.1, 0.3, [[("Numbers match the superscripts on each slide. After each source we state how "
                                          "deeply it was read; nothing is labelled verified beyond that.", False)]],
            9, MUTED)
    textbox(sl, 0.62, 1.50, 12.1, 5.2, [[(f"{i}. ", True, NAVY), (r, False)] for i, r in enumerate(REFS, 1)],
            9.8, BODY, space=5)
    notes(sl, "Full list with verification notes: 2_research/literature.csv and research_matrix.csv (41 + 16 entries).")


IMAGES = [
    ("fig_problem.png", "Problem illustration, drawn to scale", "make_figures.py · strip_v0_1_geometry.json"),
    ("fig_journey_before.png", "Before-state pain-point map", "make_figures.py · cited sources 1–3"),
    ("fig_journey_after.png", "After-state intended workflow", "make_figures.py · firmware states"),
    ("fig_block.png", "System block diagram", "make_figures.py · verification logs"),
    ("fig_dataflow.png", "Firmware scan and state logic", "make_figures.py · reader_v0_1 source"),
    ("fig_cad.png", "SolidWorks renders + drawing", "render_cad.py · 5_cad/solidworks (our models)"),
    ("fig_strip.png", "Strip and E2 coupon layers", "make_figures.py · our DXF files"),
    ("fig_schematic.png", "KiCad schematic export", "kicad-cli · reader_v0_1.kicad_sch (ours)"),
    ("fig_window.png", "Measurement window", "make_figures.py · calc_readout.py"),
    ("fig_maker.png", "SolidWorks screenshot + logs", "screen capture of our machine · verbatim tool output"),
    ("fig_decision.png", "Decision matrix + iterations", "make_figures.py · decision_matrix.md"),
    ("fig_fmea.png", "FMEA risk ranking", "make_figures.py · fmea_v0_1.csv"),
    ("fig_roadmap.png", "4-month roadmap", "make_figures.py · E2/bench/E4 plans"),
    ("fig_challenges.png", "Seal clash + lead-break diagrams", "make_figures.py · geometry + FMEA"),
]


def image_sources(sl):
    delete(sl, *range(523, 550), 550)
    pill(sl, "ALL IMAGES ORIGINAL", "green")
    textbox(sl, 0.58, 1.02, 12.17, 0.45, [[("Every image in this deck was produced by the team from our own project files. "
                                            "No stock photos, no AI-generated images and no external images are used.",
                                            False)]], 9.5, BODY)
    for k, (f, what, src) in enumerate(IMAGES):
        col, row = k % 2, k // 2
        x, y = 0.58 + col * 6.15, 1.55 + row * 0.72
        picture_in(sl, FIG / f, x, y, 0.94, 0.58)
        textbox(sl, x + 1.05, y + 0.02, 4.95, 0.6, [[(what, True, NAVY)], [("Source: ", False, MUTED), (src, False, BLUE)],
                                                    [(f"1_presentation/figures/{f}", False, MUTED)]], 7.6, space=0)
    notes(sl, "Generators: 1_presentation/render_cad.py and make_figures.py.")


# ================================================================== build
def build():
    prs = Presentation(TEMPLATE)
    t = list(prs.slides)
    e4, e6, e8, e10, e12, e15, e18 = t[3], t[5], t[7], t[9], t[11], t[14], t[17]
    s_before, s_after = duplicate(prs, e4), duplicate(prs, e4)          # duplicates first, from pristine examples
    s_fw, s_strip = duplicate(prs, e8), duplicate(prs, e8)
    s_bom = duplicate(prs, e10)
    s_iter = duplicate(prs, e12)
    s_road = duplicate(prs, e15)
    s_sol = duplicate(prs, e18)
    order = [t[0], e4, s_before, s_after, e6, s_fw, e8, s_strip, e10, s_bom, e12, s_iter, e15, s_road, e18, s_sol,
             t[18], t[19]]
    cover(t[0])
    problem(e4)
    journey_before(s_before)
    journey_after(s_after)
    arch(e6)
    firmware(s_fw)
    cad(e8)
    strip(s_strip)
    schematic(e10)
    bom(s_bom)
    maker(e12)
    iterations(s_iter)
    safety(e15)
    roadmap(s_road)
    challenges(e18)
    solutions(s_sol)
    references(t[18])
    image_sources(t[19])
    reorder_and_prune(prs, order)
    OUTDIR.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT.relative_to(ROOT)}: {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
